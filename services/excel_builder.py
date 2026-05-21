"""
Excel report builder using openpyxl.
Generates monthly reports with timesheet + invoices.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
from typing import List, Dict

class ExcelReportBuilder:
    def __init__(self, month: int, year: int):
        self.month = month
        self.year = year
        self.wb = Workbook()
        self.wb.remove(self.wb.active)  # Remove default sheet
        self.month_name = datetime(year, month, 1).strftime("%B").capitalize()

    def _get_styles(self):
        """Define reusable styles."""
        return {
            "header": Font(bold=True, size=12),
            "total": Font(bold=True, size=11),
            "total_fill": PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid"),
            "center": Alignment(horizontal="center", vertical="center"),
            "border": Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin")
            )
        }

    def add_timesheet_sheet(self, entries: List[Dict]):
        """Add timesheet sheet with entries."""
        ws = self.wb.create_sheet("Timesheet")
        styles = self._get_styles()

        # Headers
        headers = ["Date", "JJ-mois", "Client", "Affaires/Dossier", "H. Normales", "H. Sup", "Remarques"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = styles["header"]
            cell.alignment = styles["center"]
            cell.border = styles["border"]

        # Data rows
        row = 2
        total_h_normales = 0
        total_h_sup = 0

        for entry in entries:
            ws.cell(row=row, column=1).value = entry.get("date")
            ws.cell(row=row, column=2).value = entry.get("jj_mois")
            ws.cell(row=row, column=3).value = entry.get("client")
            ws.cell(row=row, column=4).value = entry.get("affaire")
            ws.cell(row=row, column=5).value = entry.get("h_normales", 0)
            ws.cell(row=row, column=6).value = entry.get("h_sup", 0)
            ws.cell(row=row, column=7).value = entry.get("remarques", "")

            total_h_normales += entry.get("h_normales", 0)
            total_h_sup += entry.get("h_sup", 0)

            # Apply borders
            for col in range(1, 8):
                ws.cell(row=row, column=col).border = styles["border"]

            row += 1

        # Total row
        total_row = row
        ws.cell(row=total_row, column=1).value = "TOTAL"
        ws.cell(row=total_row, column=1).font = styles["total"]

        ws.cell(row=total_row, column=5).value = total_h_normales
        ws.cell(row=total_row, column=5).font = styles["total"]
        ws.cell(row=total_row, column=5).fill = styles["total_fill"]

        ws.cell(row=total_row, column=6).value = total_h_sup
        ws.cell(row=total_row, column=6).font = styles["total"]
        ws.cell(row=total_row, column=6).fill = styles["total_fill"]

        # Column widths
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 10
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 20
        ws.column_dimensions["E"].width = 12
        ws.column_dimensions["F"].width = 10
        ws.column_dimensions["G"].width = 20

        return total_h_normales, total_h_sup

    def add_invoices_sheet(self, invoices: List[Dict]):
        """Add invoices sheet with email metadata + PDFs."""
        ws = self.wb.create_sheet("Factures")
        styles = self._get_styles()

        # Headers
        headers = ["Date email", "Expéditeur", "Objet", "Fichier PDF", "Lien Dropbox"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = styles["header"]
            cell.alignment = styles["center"]
            cell.border = styles["border"]

        # Data rows
        row = 2
        for invoice in invoices:
            received_date = invoice.get("received_date", "")
            if received_date:
                # Parse ISO datetime if needed
                try:
                    dt = datetime.fromisoformat(received_date.replace("Z", "+00:00"))
                    received_date = dt.strftime("%Y-%m-%d")
                except:
                    pass

            ws.cell(row=row, column=1).value = received_date
            ws.cell(row=row, column=2).value = invoice.get("from", "")
            ws.cell(row=row, column=3).value = invoice.get("subject", "")

            # PDFs (concatenated if multiple)
            pdf_names = [pdf.get("name", "") for pdf in invoice.get("pdf_attachments", [])]
            ws.cell(row=row, column=4).value = "; ".join(pdf_names)

            ws.cell(row=row, column=5).value = invoice.get("dropbox_link", "")

            # Apply borders
            for col in range(1, 6):
                ws.cell(row=row, column=col).border = styles["border"]

            row += 1

        # Column widths
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 25
        ws.column_dimensions["E"].width = 30

    def add_summary_sheet(self, total_h_normales: float, total_h_sup: float, nb_invoices: int):
        """Add summary sheet."""
        ws = self.wb.create_sheet("Résumé", 0)  # First sheet
        styles = self._get_styles()

        # Title
        ws.cell(row=1, column=1).value = f"Rapport {self.month_name} {self.year}"
        ws.cell(row=1, column=1).font = Font(bold=True, size=14)

        # Summary data
        data = [
            ("Mois", f"{self.month_name} {self.year}"),
            ("Total H. Normales", total_h_normales),
            ("Total H. Sup", total_h_sup),
            ("Nb factures trouvées", nb_invoices),
            ("Généré le", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        ]

        row = 3
        for label, value in data:
            ws.cell(row=row, column=1).value = label
            ws.cell(row=row, column=1).font = Font(bold=True)
            ws.cell(row=row, column=2).value = value
            row += 1

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 20

    def save(self, filepath: str) -> bool:
        """Save the workbook to file."""
        try:
            self.wb.save(filepath)
            return True
        except Exception as e:
            print(f"Error saving Excel file: {e}")
            return False

    def create_report(self, entries: List[Dict], invoices: List[Dict]) -> str:
        """Create a complete report and return the summary."""
        h_normales, h_sup = self.add_timesheet_sheet(entries)
        self.add_invoices_sheet(invoices)
        self.add_summary_sheet(h_normales, h_sup, len(invoices))

        summary = f"""
**Rapport {self.month_name} {self.year}**
- Jours travaillés: {len(entries)}
- Total H. Normales: {h_normales}h
- Total H. Sup: {h_sup}h
- Factures trouvées: {len(invoices)}
        """.strip()

        return summary
