"""
Update timesheet Excel file (Feuille d'heures Mai 2026.xlsx, etc.) directly.
Sheet name: FH
Structure : header row 8, data rows 9-41
Columns: A=DATES | B=CLIENTS | C=AFFAIRES | D=H.Normales | E=H.Sup | F=OBSERVATIONS
"""

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from datetime import datetime, date
import os

FRENCH_MONTHS = {
    1: "Janvier", 2: "Février", 3: "Mars", 4: "Avril", 5: "Mai", 6: "Juin",
    7: "Juillet", 8: "Août", 9: "Septembre", 10: "Octobre", 11: "Novembre", 12: "Décembre"
}

ABSENCE_KEYWORDS = ["congé", "conge", "maladie", "rtt", "férié", "ferie", "absent", "vacances"]

LOGO_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "../assets/logo_cadindus.jpeg")


def get_excel_path(base_path: str, target_date: date = None) -> str:
    if target_date is None:
        target_date = date.today()
    month_name = FRENCH_MONTHS[target_date.month]
    return os.path.join(base_path, f"Feuille d'heures {month_name} {target_date.year}.xlsx")


def get_h_normales(target_date: date) -> int:
    """Lun–Jeu = 8h, Ven = 7h"""
    return 7 if target_date.weekday() == 4 else 8


def parse_response(text: str, target_date: date) -> dict:
    """
    Parse user Discord message into a timesheet entry dict.

    Absence: "Congé" / "Maladie" / "RTT" / "Férié - Nom"
    Work:    "Client, Affaire" or "Client | Affaire"
    Work+:   "Client, Affaire, H.Sup" or "Client | Affaire | H.Sup"
    Work++:  "Client, Affaire, H.Sup, Remarque"
    """
    text = text.strip()

    if any(text.lower().startswith(kw) for kw in ABSENCE_KEYWORDS):
        return {
            "client": None,
            "affaire": text,
            "h_normales": 0,
            "h_sup": 0,
            "observations": None
        }

    # Accept pipe (|), comma (,), or slash (/) as separators
    if "|" in text:
        parts = [p.strip() for p in text.split("|")]
    elif "/" in text:
        parts = [p.strip() for p in text.split("/")]
    else:
        parts = [p.strip() for p in text.split(",")]
    h_sup = 0
    if len(parts) > 2:
        if parts[2] == "":
            h_sup = "__IDEM__"
        else:
            try:
                h_sup = float(parts[2].replace("h", "").replace("H", "").strip())
            except ValueError:
                h_sup = 0

    affaire = parts[1] if len(parts) > 1 and parts[1] else ("__IDEM__" if len(parts) > 1 else parts[0])

    return {
        "client": parts[0] if parts[0] else "__IDEM__",
        "affaire": affaire,
        "h_normales": get_h_normales(target_date),
        "h_sup": h_sup,
        "observations": parts[3] if len(parts) > 3 else None
    }


def _add_header(ws, month: int, year: int):
    """En-tête style Numbers : titre A1:D6 à gauche, logo E1:F6 à droite."""
    ws._images = []

    # Démerger anciens headers possibles
    for mr in list(ws.merged_cells.ranges):
        coord = str(mr)
        if coord in ("A1:F6", "A1:D6", "A1:D2", "A3:D4", "A3:B4", "C3:D4", "A5:D6", "E1:F6", "A1:D10", "E1:I10", "A5:I10"):
            try:
                ws.unmerge_cells(coord)
            except Exception:
                pass

    month_name = FRENCH_MONTHS[month].upper()

    # Ligne 1-2 : "FEUILLE D'HEURES"
    ws.merge_cells("A1:D2")
    ws["A1"].value = "FEUILLE D'HEURES"
    ws["A1"].font = Font(name="Calibri", size=11)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    # Ligne 3-4 : "MOIS DE" (cellule normale) + "MAI 2026" (cellule bold)
    # Deux cellules séparées = seul moyen fiable pour partial bold dans Numbers
    ws.merge_cells("A3:B4")
    ws["A3"].value = "MOIS DE"
    ws["A3"].font = Font(name="Calibri", size=11)
    ws["A3"].alignment = Alignment(horizontal="right", vertical="center")
    ws.merge_cells("C3:D4")
    ws["C3"].value = f"{month_name} {year}"
    ws["C3"].font = Font(name="Calibri", size=14, bold=True)
    ws["C3"].alignment = Alignment(horizontal="left", vertical="center")

    # Ligne 5-6 : "Jérôme SCHNEIDER"
    ws.merge_cells("A5:D6")
    ws["A5"].value = "Jérôme SCHNEIDER"
    ws["A5"].font = Font(name="Calibri", size=11)
    ws["A5"].alignment = Alignment(horizontal="center", vertical="center")

    # Zone logo à droite E1:F6
    ws.merge_cells("E1:F6")
    ws["E1"].alignment = Alignment(horizontal="center", vertical="center")

    # Logo centré dans E1:F6
    # E width = 13 chars × 7px + 5 ≈ 96 px, F width = 35 chars × 7px + 5 ≈ 250 px → total ≈ 346 px
    # Logo 250x82 px → marge horizontale = (346-250)/2 ≈ 48 px depuis E1
    # Hauteur ligne 1-6 ≈ 6 × 15.75 ≈ 94 px → marge verticale = (94-82)/2 ≈ 6 px
    if os.path.exists(LOGO_PATH):
        try:
            img = XLImage(LOGO_PATH)
            marker = AnchorMarker(
                col=4, colOff=pixels_to_EMU(48),
                row=0, rowOff=pixels_to_EMU(6),
            )
            size = XDRPositiveSize2D(cx=pixels_to_EMU(250), cy=pixels_to_EMU(82))
            img.anchor = OneCellAnchor(_from=marker, ext=size)
            ws.add_image(img)
        except Exception as e:
            print(f"[excel_updater] logo add failed: {e}", flush=True)

    # Print A4 portrait
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.3
    ws.page_margins.bottom = 0.3
    ws.print_options.horizontalCentered = True
    ws.print_area = "A1:F45"


def _autofit_columns(ws):
    """Mise en forme tableau style Numbers : header row 8, data 9-41, 6 colonnes A-F contiguës."""
    from openpyxl.utils import get_column_letter

    # Largeurs identiques au Numbers du user
    WIDTHS = {"A": 14, "B": 22, "C": 30, "D": 10, "E": 13, "F": 35}
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
        ws.column_dimensions[col].hidden = False

    # Cacher toutes les colonnes au-delà de F
    for col_idx in range(7, 60):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].hidden = True
        ws.column_dimensions[letter].width = 0

    medium = Side(style="medium", color="000000")
    thin   = Side(style="thin",   color="000000")

    # === Header row 8 ===
    for col_letter in "ABCDEF":
        cell = ws[f"{col_letter}8"]
        cell.font = Font(name="Calibri", size=10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left  = medium if col_letter == "A" else thin
        right = medium if col_letter == "F" else thin
        cell.border = Border(left=left, right=right, top=medium, bottom=medium)
    ws.row_dimensions[8].height = 30

    # === Data rows 9-41 ===
    for r in range(9, 42):
        ws.row_dimensions[r].height = 18
        for col_letter in "ABCDEF":
            cell = ws[f"{col_letter}{r}"]
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            left   = medium if col_letter == "A" else thin
            right  = medium if col_letter == "F" else thin
            bottom = medium if r == 41 else thin
            cell.border = Border(left=left, right=right, top=thin, bottom=bottom)
            if col_letter == "A":
                cell.font = Font(name="Calibri", size=13, bold=True)
                if isinstance(cell.value, datetime):
                    cell.number_format = "d-mmm"
            else:
                # Auto-shrink police selon longueur
                size = 10
                if cell.value is not None and not isinstance(cell.value, (int, float)):
                    text_len = len(str(cell.value))
                    eff_width = WIDTHS.get(col_letter, 15)
                    ratio = text_len / eff_width
                    if ratio > 1.5:
                        size = 8
                    elif ratio > 1.0:
                        size = 9
                cell.font = Font(name="Calibri", size=size)

    # === Footer (lignes 42 = total, 44 = date) ===
    ws["B44"].number_format = "dd/mm/yyyy"


def update_excel_entry(base_path: str, entry: dict, target_date: date = None) -> str:
    """
    Find the row for target_date in the Excel file and write the entry.
    On VPS (file not present locally), downloads from Dropbox and re-uploads after save.
    Returns a Discord-ready status message.
    """
    if target_date is None:
        target_date = date.today()

    excel_path = get_excel_path(base_path, target_date)
    # VPS mode: enabled when DROPBOX_SYNC env var is true (no local Dropbox sync available)
    dropbox_sync = os.getenv("DROPBOX_SYNC", "").lower() in ("1", "true", "yes")

    if dropbox_sync:
        # Always download fresh from Dropbox before editing
        try:
            from services.dropbox_client import DropboxClient
            from config import DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN, DROPBOX_VAULT_PATH
            dbx = DropboxClient(DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN)
            month_name = FRENCH_MONTHS[target_date.month]
            dropbox_path = f"{DROPBOX_VAULT_PATH}/Feuille d'heures {month_name} {target_date.year}.xlsx"
            os.makedirs(base_path, exist_ok=True)
            if not dbx.download_binary(dropbox_path, excel_path):
                return f"❌ Fichier non trouvé: {os.path.basename(excel_path)}"
        except Exception as e:
            return f"❌ Erreur Dropbox: {e}"
    elif not os.path.exists(excel_path):
        return f"❌ Fichier non trouvé: {os.path.basename(excel_path)}"

    wb = load_workbook(excel_path)
    if "FH" not in wb.sheetnames:
        return f"❌ Feuille 'FH' introuvable dans {os.path.basename(excel_path)}"

    ws = wb["FH"]

    target_row = None
    for row in ws.iter_rows(min_row=9, max_row=41):
        cell_val = row[0].value
        if isinstance(cell_val, datetime) and cell_val.date() == target_date:
            target_row = row[0].row
            break

    if target_row is None:
        return f"❌ Date {target_date.strftime('%d/%m')} non trouvée dans {os.path.basename(excel_path)}"

    def _get_prev(col_letter):
        val = ws[f"{col_letter}{target_row - 1}"].value
        return val if val is not None and str(val).strip() else None

    if entry["client"] == "__IDEM__":
        entry["client"] = _get_prev("B")

    if entry["affaire"] == "__IDEM__":
        entry["affaire"] = _get_prev("C") or ""

    if entry["h_sup"] == "__IDEM__":
        entry["h_sup"] = _get_prev("E") or 0

    # Nouvelles colonnes : B=client, C=affaire, D=h_normales, E=h_sup, F=observations
    ws[f"B{target_row}"] = entry["client"]
    ws[f"C{target_row}"] = entry["affaire"]
    ws[f"D{target_row}"] = entry["h_normales"]
    ws[f"E{target_row}"] = entry["h_sup"]
    ws[f"F{target_row}"] = entry["observations"]

    _add_header(ws, target_date.month, target_date.year)
    _autofit_columns(ws)
    wb.save(excel_path)

    # VPS mode: re-upload to Dropbox after save
    if dropbox_sync:
        try:
            from services.dropbox_client import DropboxClient
            from config import DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN, DROPBOX_VAULT_PATH
            dbx = DropboxClient(DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN)
            month_name = FRENCH_MONTHS[target_date.month]
            dropbox_path = f"{DROPBOX_VAULT_PATH}/Feuille d'heures {month_name} {target_date.year}.xlsx"
            dbx.upload_binary_file(excel_path, dropbox_path)
        except Exception as e:
            print(f"[excel_updater] Dropbox re-upload failed: {e}", flush=True)

    day_label = target_date.strftime("%d/%m")
    affaire = entry["affaire"] or ""
    h_norm = entry["h_normales"]
    h_sup = entry["h_sup"]
    client_info = f" ({entry['client']})" if entry["client"] else ""
    return f"✓ {day_label}{client_info} — {affaire} — {h_norm}h normales, {h_sup}h sup"


def create_month_template(base_path: str, year: int, month: int) -> str:
    """Create a new monthly timesheet Excel file with workdays pre-filled in column A
    and French/Alsace-Moselle holidays pre-marked in column C."""
    from calendar import monthrange
    from services.holidays import french_holidays
    target_date = date(year, month, 1)
    excel_path = get_excel_path(base_path, target_date)
    dropbox_sync = os.getenv("DROPBOX_SYNC", "").lower() in ("1", "true", "yes")
    month_name = FRENCH_MONTHS[month]

    # Skip if file already exists in Dropbox
    if dropbox_sync:
        try:
            from services.dropbox_client import DropboxClient
            from config import DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN, DROPBOX_VAULT_PATH
            dbx = DropboxClient(DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN)
            dropbox_path = f"{DROPBOX_VAULT_PATH}/Feuille d'heures {month_name} {year}.xlsx"
            os.makedirs(base_path, exist_ok=True)
            try:
                dbx.dbx.files_get_metadata(dropbox_path)
                return f"ℹ️ {month_name} {year} existe déjà"
            except Exception:
                pass
        except Exception:
            pass

    # Use openpyxl to create from scratch using same structure as existing template
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "FH"

    # Header (rows 1-6) + table header (row 8)
    _add_header(ws, month, year)
    ws["A8"] = "DATES"
    ws["B8"] = "CLIENTS"
    ws["C8"] = "AFFAIRES  ou  DOSSIERS"
    ws["D8"] = "H. Normales"
    ws["E8"] = "H. Sup"
    ws["F8"] = "Observations"

    # Fill column A with workdays only (Mon–Fri)
    holidays = french_holidays(year, alsace_moselle=True)
    days_in_month = monthrange(year, month)[1]
    row = 9
    for d in range(1, days_in_month + 1):
        the_date = date(year, month, d)
        if the_date.weekday() >= 5:  # skip weekends
            continue
        ws[f"A{row}"] = datetime(year, month, d)
        if the_date in holidays:
            ws[f"C{row}"] = f"Férié - {holidays[the_date]}"
            ws[f"D{row}"] = 0
        row += 1

    _autofit_columns(ws)
    wb.save(excel_path)

    if dropbox_sync:
        try:
            from services.dropbox_client import DropboxClient
            from config import DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN, DROPBOX_VAULT_PATH
            dbx = DropboxClient(DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN)
            dropbox_path = f"{DROPBOX_VAULT_PATH}/Feuille d'heures {month_name} {year}.xlsx"
            dbx.upload_binary_file(excel_path, dropbox_path)
        except Exception as e:
            return f"❌ Upload échoué: {e}"

    return f"✓ Créé : Feuille d'heures {month_name} {year}.xlsx ({row - 9} jours ouvrés)"


def clear_row(base_path: str, target_date: date) -> str:
    """Clear B/C/D/E/F for a given date row. Used by /annuler."""
    excel_path = get_excel_path(base_path, target_date)
    dropbox_sync = os.getenv("DROPBOX_SYNC", "").lower() in ("1", "true", "yes")
    month_name = FRENCH_MONTHS[target_date.month]

    if dropbox_sync:
        try:
            from services.dropbox_client import DropboxClient
            from config import DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN, DROPBOX_VAULT_PATH
            dbx = DropboxClient(DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN)
            dropbox_path = f"{DROPBOX_VAULT_PATH}/Feuille d'heures {month_name} {target_date.year}.xlsx"
            os.makedirs(base_path, exist_ok=True)
            if not dbx.download_binary(dropbox_path, excel_path):
                return f"❌ Fichier non trouvé: {os.path.basename(excel_path)}"
        except Exception as e:
            return f"❌ Erreur Dropbox: {e}"
    elif not os.path.exists(excel_path):
        return f"❌ Fichier non trouvé: {os.path.basename(excel_path)}"

    wb = load_workbook(excel_path)
    ws = wb["FH"]
    target_row = None
    for row in ws.iter_rows(min_row=9, max_row=41):
        if isinstance(row[0].value, datetime) and row[0].value.date() == target_date:
            target_row = row[0].row
            break
    if target_row is None:
        return f"❌ Date {target_date.strftime('%d/%m')} non trouvée"

    for col in "BCDEF":
        ws[f"{col}{target_row}"] = None
    _autofit_columns(ws)
    wb.save(excel_path)

    if dropbox_sync:
        try:
            from services.dropbox_client import DropboxClient
            from config import DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN, DROPBOX_VAULT_PATH
            dbx = DropboxClient(DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN)
            dropbox_path = f"{DROPBOX_VAULT_PATH}/Feuille d'heures {month_name} {target_date.year}.xlsx"
            dbx.upload_binary_file(excel_path, dropbox_path)
        except Exception as e:
            print(f"[excel_updater] re-upload failed: {e}", flush=True)
    return f"✓ {target_date.strftime('%d/%m')} effacé"


def summarize_month(base_path: str, target_date: date = None) -> str:
    """Return Discord-ready summary of the month's timesheet."""
    if target_date is None:
        target_date = date.today()
    excel_path = get_excel_path(base_path, target_date)
    month_name = FRENCH_MONTHS[target_date.month]
    dropbox_sync = os.getenv("DROPBOX_SYNC", "").lower() in ("1", "true", "yes")

    if dropbox_sync:
        try:
            from services.dropbox_client import DropboxClient
            from config import DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN, DROPBOX_VAULT_PATH
            dbx = DropboxClient(DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN)
            dropbox_path = f"{DROPBOX_VAULT_PATH}/Feuille d'heures {month_name} {target_date.year}.xlsx"
            os.makedirs(base_path, exist_ok=True)
            if not dbx.download_binary(dropbox_path, excel_path):
                return f"❌ Fichier non trouvé: {os.path.basename(excel_path)}"
            # Count invoices in archive folder
            inv_count = 0
            try:
                r = dbx.dbx.files_list_folder(f"{DROPBOX_VAULT_PATH}/Feuille d'heures/{target_date.year}/{month_name} {target_date.year}")
                inv_count = sum(1 for e in r.entries if e.name.lower().endswith(".pdf"))
            except Exception:
                pass
        except Exception as e:
            return f"❌ Erreur Dropbox: {e}"
    else:
        if not os.path.exists(excel_path):
            return f"❌ Fichier non trouvé: {os.path.basename(excel_path)}"
        inv_count = 0

    wb = load_workbook(excel_path)
    ws = wb["FH"]
    filled, missing, h_norm_total, h_sup_total = [], [], 0, 0.0
    today = date.today()
    for row in ws.iter_rows(min_row=9, max_row=41):
        d_cell = row[0].value
        if not isinstance(d_cell, datetime):
            continue
        d = d_cell.date()
        if d.weekday() >= 5 or d.month != target_date.month:
            continue
        has_data = row[1].value or row[2].value
        if has_data:
            try:
                h_norm_total += int(row[3].value or 0)
            except (TypeError, ValueError):
                pass
            try:
                h_sup_total += float(row[4].value or 0)
            except (TypeError, ValueError):
                pass
            filled.append(d)
        elif d <= today:
            missing.append(d)

    miss_str = ", ".join(d.strftime("%d/%m") for d in missing) if missing else "aucun ✓"
    return (
        f"📊 **{month_name} {target_date.year}**\n"
        f"Jours saisis : **{len(filled)}** · H.Normales : **{h_norm_total}h** · H.Sup : **{h_sup_total}h**\n"
        f"Jours ouvrés manquants (≤ aujourd'hui) : {miss_str}\n"
        f"Factures archivées : **{inv_count}** 📄"
    )


def update_observation(base_path: str, vendor: str, total: str, target_date: date) -> str:
    """Append 'Vendor : total€' in column F (Observations) of the target date row."""
    excel_path = get_excel_path(base_path, target_date)
    dropbox_sync = os.getenv("DROPBOX_SYNC", "").lower() in ("1", "true", "yes")

    if dropbox_sync:
        try:
            from services.dropbox_client import DropboxClient
            from config import DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN, DROPBOX_VAULT_PATH
            dbx = DropboxClient(DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN)
            month_name = FRENCH_MONTHS[target_date.month]
            dropbox_path = f"{DROPBOX_VAULT_PATH}/Feuille d'heures {month_name} {target_date.year}.xlsx"
            os.makedirs(base_path, exist_ok=True)
            if not dbx.download_binary(dropbox_path, excel_path):
                return f"❌ Fichier non trouvé: {os.path.basename(excel_path)}"
        except Exception as e:
            return f"❌ Erreur Dropbox: {e}"
    elif not os.path.exists(excel_path):
        return f"❌ Fichier non trouvé: {os.path.basename(excel_path)}"

    wb = load_workbook(excel_path)
    if "FH" not in wb.sheetnames:
        return f"❌ Feuille 'FH' introuvable dans {os.path.basename(excel_path)}"

    ws = wb["FH"]
    target_row = None
    for row in ws.iter_rows(min_row=9, max_row=41):
        cell_val = row[0].value
        if isinstance(cell_val, datetime) and cell_val.date() == target_date:
            target_row = row[0].row
            break

    if target_row is None:
        return f"❌ Date {target_date.strftime('%d/%m')} non trouvée dans {os.path.basename(excel_path)}"

    existing = ws[f"F{target_row}"].value or ""
    separator = " | " if existing else ""
    ws[f"F{target_row}"] = f"{existing}{separator}{vendor} : {total}€"

    _autofit_columns(ws)
    wb.save(excel_path)

    if dropbox_sync:
        try:
            from services.dropbox_client import DropboxClient
            from config import DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN, DROPBOX_VAULT_PATH
            dbx = DropboxClient(DROPBOX_APP_KEY, DROPBOX_APP_SECRET, DROPBOX_REFRESH_TOKEN)
            month_name = FRENCH_MONTHS[target_date.month]
            dropbox_path = f"{DROPBOX_VAULT_PATH}/Feuille d'heures {month_name} {target_date.year}.xlsx"
            dbx.upload_binary_file(excel_path, dropbox_path)
        except Exception as e:
            print(f"[excel_updater] Dropbox re-upload failed: {e}", flush=True)

    return f"✓ {target_date.strftime('%d/%m')} — {vendor} : {total}€ ajouté en observations"
